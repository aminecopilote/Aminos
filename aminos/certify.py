"""Certified translation: declaration, numbering, HMAC code/seal, QR, one-page layout, evaluation gate.

Process (see README): build the .docx WITHOUT the certification block -> CODE = HMAC over the body text
-> QR -> rebuild WITH the block (QR left, mentions right) -> verify -> seal (only after the translator's
validation of the evaluation sheet). The HMAC key never leaves the machine that holds the key file.

Hash scope: only body-level paragraphs up to the "FIN DE TRADUCTION" line, so adding the declaration /
certification block does not change the text fingerprint.
"""
from __future__ import annotations

import csv
import datetime as dt
import hashlib
import hmac
import io
import re
import secrets
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree as ET

from . import prompts

SIGNER_LATIN = "Alami"
SIGNER_AR = "علمي"
CITY = {"fr": "Rabat", "en": "Rabat", "ar": "الرباط"}
TITLE = {"fr": "traducteur certifié", "en": "certified translator", "ar": "مترجم معتمد"}
BANNER = {"fr": "TRADUCTION CERTIFIÉE CONFORME", "en": "CERTIFIED TRUE TRANSLATION",
          "ar": "ترجمة معتمدة مطابقة للأصل"}
END_LINE = {"fr": "FIN DE TRADUCTION", "en": "END OF TRANSLATION", "ar": "نهاية الترجمة"}
END_MARK = {l: f"------------------------{t}---------------------------" for l, t in END_LINE.items()}
LABELS = {"fr": ("Référence", "Date", "Code de vérification"),
          "en": ("Reference", "Date", "Verification code"),
          "ar": ("المرجع", "التاريخ", "رمز التحقق")}
LANG_ALIASES = {"fr": "fr", "français": "fr", "francais": "fr", "french": "fr",
                "ar": "ar", "arabe": "ar", "arabic": "ar", "العربية": "ar",
                "en": "en", "anglais": "en", "english": "en"}
LANG_NAMES = {  # name of each language, written in the language of the declaration
    "fr": {"fr": "français", "ar": "arabe", "en": "anglais"},
    "en": {"fr": "French", "ar": "Arabic", "en": "English"},
    "ar": {"fr": "الفرنسية", "ar": "العربية", "en": "الإنجليزية"}}
MONTHS = {
    "fr": "janvier février mars avril mai juin juillet août septembre octobre novembre décembre".split(),
    "en": "January February March April May June July August September October November December".split(),
    "ar": "يناير فبراير مارس أبريل ماي يونيو يوليوز غشت شتنبر أكتوبر نونبر دجنبر".split()}
CRITERIA = ["fidelite", "terminologie", "conformite_marocaine", "structure", "ponctuation", "formats"]
CRITERIA_LABEL = {"fidelite": "Fidélité", "terminologie": "Terminologie & précision",
                  "conformite_marocaine": "Conformité terminologique marocaine",
                  "structure": "Structure & mise en page", "ponctuation": "Ponctuation & style",
                  "formats": "Formats & conventions"}
LRI, PDI = "⁦", "⁩"
_REF = re.compile(r"^\d{5}\.\d{4}$")


def lang_code(name: str) -> str:
    try:
        return LANG_ALIASES[name.strip().lower()]
    except KeyError:
        raise ValueError(f"Langue non gérée : {name} (fr, ar, en)") from None


def long_date(date: str, lang: str) -> str:
    d = dt.date.fromisoformat(date)
    m = MONTHS[lang][d.month - 1]
    return {"fr": f"{d.day}{'er' if d.day == 1 else ''} {m} {d.year}", "en": f"{m} {d.day}, {d.year}",
            "ar": f"{d.day} {m} {d.year}"}[lang]


def declaration(source: str, target: str, ref: str | None = None, date: str | None = None) -> str:
    """Faithfulness declaration, written only in the language of the translation."""
    s, t = lang_code(source), lang_code(target)
    date = date or dt.date.today().isoformat()
    n, d, city = LANG_NAMES[t], long_date(date, t), CITY[t]
    refpart = {"fr": f" Référence : {ref}.", "en": f" Reference: {ref}.", "ar": f" المرجع: {ref}."}[t] if ref else ""
    if t == "fr":
        de = {"fr": "du français", "ar": "de l'arabe", "en": "de l'anglais"}[s]
        vers = {"fr": "le français", "ar": "l'arabe", "en": "l'anglais"}[t]
        return (f"Je soussigné, {SIGNER_LATIN}, {TITLE[t]}, certifie que la présente traduction {de} vers "
                f"{vers} est fidèle et conforme au document original qui m'a été présenté.{refpart} "
                f"Fait à {city}, le {d}.")
    if t == "en":
        return (f"I, the undersigned, {SIGNER_LATIN}, {TITLE[t]}, certify that this translation from {n[s]} into "
                f"{n[t]} is a faithful and true rendering of the original document presented to me.{refpart} "
                f"Done at {city}, on {d}.")
    return (f"أنا الموقع أدناه {SIGNER_AR}، {TITLE[t]}، أشهد بأن هذه الترجمة من اللغة {n[s]} إلى اللغة {n[t]} "
            f"مطابقة للوثيقة الأصلية التي عُرضت علي.{refpart} حُرر بـ{city} في {d}.")


# ---- registry & numbering ---------------------------------------------------
REGISTRY_FIELDS = ["ref", "date", "code", "seal", "status"]


def read_registry(path: str | Path) -> list[dict]:
    p = Path(path)
    if not p.exists():
        return []
    with p.open(encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def write_registry(path: str | Path, rows: list[dict]) -> None:
    with Path(path).open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, REGISTRY_FIELDS)
        w.writeheader()
        w.writerows({k: r.get(k, "") for k in REGISTRY_FIELDS} for r in rows)


def next_ref(registry: str | Path, year: int | None = None) -> str:
    """NNNNN.AAAA, incremented by 1 over the highest number in the registry (all years)."""
    year = year or dt.date.today().year
    nums = [int(r["ref"].split(".")[0]) for r in read_registry(registry) if _REF.match(r.get("ref", ""))]
    return f"{(max(nums) + 1) if nums else 1:05d}.{year}"


def upsert(registry: str | Path, **fields) -> None:
    rows = read_registry(registry)
    for r in rows:
        if r["ref"] == fields["ref"]:
            r.update({k: v for k, v in fields.items() if v is not None})
            break
    else:
        rows.append(fields)
    write_registry(registry, rows)


# ---- keys, fingerprints ---------------------------------------------------------
def make_key(path: str | Path) -> None:
    p = Path(path)
    if p.exists():
        raise FileExistsError(f"{p} existe déjà : une clé ne doit jamais être écrasée")
    p.write_text(secrets.token_hex(32), encoding="utf-8")
    try:
        p.chmod(0o600)
    except OSError:
        pass


def load_key(path: str | Path) -> bytes:
    k = Path(path).read_text(encoding="utf-8").strip()
    if len(k) < 32:
        raise ValueError("clé HMAC trop courte")
    return k.encode()


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def body_text(docx_path: str | Path) -> str:
    """Body-level paragraphs up to (excluding) the END line; used for the CODE."""
    with zipfile.ZipFile(docx_path) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    out = []
    for p in root.find(f"{_W}body").findall(f"{_W}p"):
        text = "".join(t.text or "" for t in p.iter(f"{_W}t"))
        if any(m in text for m in END_LINE.values()):
            break
        out.append(text)
    return "\n".join(out)


def _mac(key: bytes, ref: str, date: str, digest: str) -> str:
    return hmac.new(key, f"{ref}|{date}|{SIGNER_LATIN}|{digest}".encode(), hashlib.sha256).hexdigest().upper()


def compute_code(docx_path, key: bytes, ref: str, date: str) -> str:
    return _mac(key, ref, date, hashlib.sha256(body_text(docx_path).encode()).hexdigest())


def compute_seal(docx_path, key: bytes, ref: str, date: str) -> str:
    return _mac(key, ref, date, hashlib.sha256(Path(docx_path).read_bytes()).hexdigest())


def qr_payload(ref: str, date: str, code: str) -> str:
    return f"REF={ref};DATE={date};CODE={code}"


def make_qr_png(payload: str, size_mm: float = 28, dpi: int = 300) -> bytes:
    from PIL import Image
    from reportlab.graphics.barcode import qr
    w = qr.QrCodeWidget(payload, barLevel="M")
    w.qr.make()
    n = w.qr.getModuleCount()
    px = max(int(size_mm / 25.4 * dpi) // (n + 8), 4)  # 4-module quiet zone each side
    img = Image.new("L", ((n + 8) * px,) * 2, 255)
    for y in range(n):
        for x in range(n):
            if w.qr.isDark(y, x):
                img.paste(0, ((x + 4) * px, (y + 4) * px, (x + 5) * px, (y + 5) * px))
    buf = io.BytesIO()
    img.save(buf, "PNG", dpi=(dpi, dpi))
    return buf.getvalue()


def read_qr(png: bytes) -> str:
    import cv2
    import numpy as np
    img = cv2.imdecode(np.frombuffer(png, np.uint8), cv2.IMREAD_GRAYSCALE)
    text, _, _ = cv2.QRCodeDetector().detectAndDecode(img)
    return text


# ---- layout -------------------------------------------------------------------
def isolate_digits(text: str) -> str:
    """Wrap tokens containing a digit in LTR isolates so dates/numbers are not reversed in Arabic."""
    return re.sub(r"[^\s()،؛:]*\d[^\s()،؛:]*", lambda m: LRI + m.group(0) + PDI, text)


def pick_font_size(paragraphs: list[str], extra_chars: int = 900) -> float:
    """Heuristic for one page: 10 pt, then 9.5, then 9. Verify with page_count()."""
    chars = sum(len(p) + 40 for p in paragraphs) + extra_chars
    return 10 if chars <= 4200 else 9.5 if chars <= 4700 else 9


def build_certified_docx(paragraphs: list[str], source: str, target: str, ref: str, date: str, path: str | Path,
                         qr_png: bytes | None = None, code: str | None = None, font_size: float | None = None,
                         font: str = "Calibri") -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    t = lang_code(target)
    rtl = t == "ar"
    size = font_size or pick_font_size(paragraphs)
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(1.27)
        s.header.is_linked_to_previous = True  # no header/footer content is ever written
    st = doc.styles["Normal"]
    st.font.name, st.font.size = font, Pt(size)
    st.element.rPr.rFonts.set(qn("w:cs"), font)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    st.paragraph_format.space_before = st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = 1.0

    def para(container, text: str, bold_all=False, center=False, before=0, ltr=False):
        p = container.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        if rtl and not ltr:
            p._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
            # In a bidi paragraph "left" is the start of the line = the right margin.
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        if rtl and not ltr:
            text = isolate_digits(text)
        for part in filter(None, re.split(r"(\*\*.+?\*\*)", text)):
            bold = bold_all or (part.startswith("**") and part.endswith("**") and len(part) > 4)
            r = p.add_run(part[2:-2] if part.startswith("**") and part.endswith("**") and len(part) > 4 else part)
            r.bold = bold or None
            r.font.name, r.font.size = font, Pt(size)
            if rtl and not ltr:
                r._r.get_or_add_rPr().append(OxmlElement("w:rtl"))
        return p

    # drop any template paragraph so the banner is the first line of the body
    for empty in list(doc.paragraphs):
        empty._p.getparent().remove(empty._p)
    para(doc, BANNER[t], bold_all=True, center=True)
    for text in paragraphs:
        para(doc, text)
    para(doc, END_MARK[t], center=True, before=7, ltr=True)
    para(doc, declaration(source, target, ref, date), before=4)

    if qr_png and code:
        l_ref, l_date, l_code = LABELS[t]
        table = doc.add_table(rows=1, cols=2)
        table.autofit = False
        left, right = table.rows[0].cells
        left.width, right.width = Cm(3.2), Cm(14.3)
        left.paragraphs[0].add_run().add_picture(io.BytesIO(qr_png), width=Cm(2.8))
        right.paragraphs[0]._p.getparent().remove(right.paragraphs[0]._p)
        for label, value in ((l_ref, ref), (l_date, long_date(date, t)), (l_code, code[:16] + "…")):
            para(right, f"{label} : {value}", ltr=False)
    doc.save(str(path))


def page_count(docx_path: str | Path) -> int | None:
    """Page count through LibreOffice (None when it is not installed)."""
    import shutil
    import subprocess
    import tempfile
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return None
    with tempfile.TemporaryDirectory() as d:
        subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", d, str(docx_path)],
                       check=True, capture_output=True, timeout=180)
        pdf = Path(d) / (Path(docx_path).stem + ".pdf")
        return len(re.findall(rb"/Type\s*/Page[^s]", pdf.read_bytes()))


# ---- evaluation gate -------------------------------------------------------------
@dataclass
class Evaluation:
    scores: dict[str, int]
    points: list[str]

    def passes(self) -> bool:
        return (set(self.scores) == set(CRITERIA) and self.scores["fidelite"] == 5
                and all(v >= 4 for k, v in self.scores.items() if k != "fidelite")
                and all(0 <= v <= 5 for v in self.scores.values()))


def write_evaluation(path: str | Path, ref: str, ev: Evaluation) -> None:
    rows = "\n".join(f"| {CRITERIA_LABEL[c]} | {ev.scores.get(c, '')} | |" for c in CRITERIA)
    pts = "\n".join(f"- {p}" for p in ev.points) or "- (aucun)"
    Path(path).write_text(
        f"# Évaluation {ref}\n\n| Critère | Claude /5 | Alami /5 |\n|---|---|---|\n{rows}\n\n"
        f"Seuil : fidélité 5/5, autres ≥ 4/5 — {'atteint' if ev.passes() else 'NON atteint : reprendre'}\n\n"
        f"## Points à confirmer sur l'original\n{pts}\n\nValidation Alami : non\n", encoding="utf-8")


def evaluation_validated(path: str | Path) -> bool:
    """True only when the sheet meets the threshold AND the translator wrote 'Validation Alami : oui'."""
    text = Path(path).read_text(encoding="utf-8")
    if not re.search(r"^Validation Alami\s*:\s*oui\s*$", text, re.M | re.I):
        return False
    scores = {}
    for c in CRITERIA:
        m = re.search(rf"\|\s*{re.escape(CRITERIA_LABEL[c])}\s*\|\s*(\d)\s*\|", text)
        if not m:
            return False
        scores[c] = int(m.group(1))
    return Evaluation(scores, []).passes()


# ---- workflow --------------------------------------------------------------------
def certify(paragraphs: list[str], source: str, target: str, key_file, registry, out_dir,
            ref: str | None = None, date: str | None = None) -> dict:
    """Steps 1-4: build, fingerprint, QR, rebuild with the block. Returns ref/code/path (not yet sealed)."""
    date = date or dt.date.today().isoformat()
    key = load_key(key_file)
    ref = ref or next_ref(registry, int(date[:4]))
    if not _REF.match(ref):
        raise ValueError("référence attendue : NNNNN.AAAA")
    if any(r["ref"] == ref and r.get("status") == "scellé" for r in read_registry(registry)):
        raise ValueError(f"{ref} est déjà scellé")
    out = Path(out_dir) / f"{ref}.docx"
    Path(out_dir).mkdir(parents=True, exist_ok=True)
    size = pick_font_size(paragraphs)
    build_certified_docx(paragraphs, source, target, ref, date, out, font_size=size)
    code = compute_code(out, key, ref, date)
    build_certified_docx(paragraphs, source, target, ref, date, out, make_qr_png(qr_payload(ref, date, code)),
                         code, size)
    upsert(registry, ref=ref, date=date, code=code, seal="", status="en attente")
    return {"ref": ref, "date": date, "code": code, "path": str(out)}


def verify(docx_path, key_file, registry, ref: str) -> bool:
    row = next((r for r in read_registry(registry) if r["ref"] == ref), None)
    if not row:
        return False
    return hmac.compare_digest(compute_code(docx_path, load_key(key_file), ref, row["date"]), row["code"])


def seal(docx_path, key_file, registry, ref: str, evaluation_file) -> str:
    if not evaluation_validated(evaluation_file):
        raise PermissionError("Aucune certification sans validation d'Alami : fiche d'évaluation non validée "
                              "ou seuil non atteint.")
    if not verify(docx_path, key_file, registry, ref):
        raise ValueError("Le code ne correspond plus au texte : refaire la certification.")
    row = next(r for r in read_registry(registry) if r["ref"] == ref)
    s = compute_seal(docx_path, load_key(key_file), ref, row["date"])
    upsert(registry, ref=ref, seal=s, status="scellé")
    return s
