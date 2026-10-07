"""Word output: Markdown-style **bold** to real bold runs, RTL for Arabic, Calibri 10 / margins 1.27 cm."""
from __future__ import annotations

import re

from . import prompts
from .certify import isolate_digits

_BOLD = re.compile(r"(\*\*.+?\*\*)")


def build_docx(paragraphs: list[str], target_lang: str, path: str, font: str = "Calibri", size: int = 10) -> None:
    try:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError as e:
        raise RuntimeError("pip install python-docx") from e

    rtl = prompts.is_rtl(target_lang)
    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(1.27)

    for text in paragraphs:
        if rtl:
            text = isolate_digits(text)
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        if rtl:
            p._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
            # in a bidi paragraph "left" is the line start, i.e. the right margin ("right" would align left)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for part in filter(None, _BOLD.split(text)):
            bold = part.startswith("**") and part.endswith("**") and len(part) > 4
            run = p.add_run(part[2:-2] if bold else part)
            run.bold = bold or None
            run.font.name, run.font.size = font, Pt(size)
            rpr = run._r.get_or_add_rPr()
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                rpr.append(fonts)
            for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                fonts.set(qn(attr), font)
            if rtl:
                rpr.append(OxmlElement("w:rtl"))
    doc.save(path)
