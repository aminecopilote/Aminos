"""Terminology engine: import glossaries, detect terms in a source text, verify usage."""
from __future__ import annotations

import csv
import json
import re
from dataclasses import dataclass, field, asdict
from pathlib import Path

from .normalize import normalize, tokens

SUPPORTED_SUFFIXES = {".csv", ".tsv", ".txt", ".json", ".xlsx", ".docx"}
_TXT_SPLIT = re.compile(r"\s*(?:\t|=|:|;|\||→|->|–|—| - )\s*")


@dataclass
class Term:
    source: str
    target: str
    domain: str = ""
    note: str = ""
    origin: str = ""

    def key(self) -> str:
        return normalize(self.source)


@dataclass
class Glossary:
    terms: dict[str, Term] = field(default_factory=dict)
    conflicts: list[tuple[Term, Term]] = field(default_factory=list)

    def __len__(self) -> int:
        return len(self.terms)

    def add(self, term: Term) -> None:
        term.source, term.target = term.source.strip(), term.target.strip()
        if not term.source or not term.target:
            return
        key = term.key()
        old = self.terms.get(key)
        if old and normalize(old.target) != normalize(term.target):
            self.conflicts.append((old, term))
            return  # first definition wins; conflict is reported
        self.terms.setdefault(key, term)

    def extend(self, terms) -> None:
        for t in terms:
            self.add(t)

    # ---- detection -------------------------------------------------
    def find(self, text: str) -> list[Term]:
        """Glossary terms present in `text` (longest match first, no overlaps)."""
        toks = tokens(text)
        by_tokens = [(tuple(tokens(t.source)), t) for t in self.terms.values()]
        by_tokens = [(tk, t) for tk, t in by_tokens if tk]
        by_tokens.sort(key=lambda p: -len(p[0]))
        used = [False] * len(toks)
        found: dict[str, Term] = {}
        for tk, term in by_tokens:
            n = len(tk)
            for i in range(len(toks) - n + 1):
                if tuple(toks[i:i + n]) == tk and not any(used[i:i + n]):
                    used[i:i + n] = [True] * n
                    found[term.key()] = term
        return list(found.values())

    def prompt_block(self, text: str) -> str:
        return "\n".join(f"- {t.source} => {t.target}" + (f" [{t.note}]" if t.note else "")
                         for t in self.find(text))

    def missing_in(self, source: str, translation: str) -> list[Term]:
        """Terms found in source whose imposed target is absent from the translation."""
        out = []
        norm_target = normalize(translation)
        for term in self.find(source):
            variants = [normalize(v) for v in re.split(r"\s*[/|]\s*", term.target) if v.strip()]
            if not any(v in norm_target for v in variants):
                out.append(term)
        return out

    # ---- persistence -----------------------------------------------
    def save(self, path: str | Path) -> None:
        Path(path).write_text(
            json.dumps([asdict(t) for t in self.terms.values()], ensure_ascii=False, indent=1),
            encoding="utf-8")

    @classmethod
    def load(cls, path: str | Path) -> "Glossary":
        g = cls()
        for row in json.loads(Path(path).read_text(encoding="utf-8")):
            g.add(Term(**row))
        return g


# ---- importers -----------------------------------------------------
_SRC_HEADERS = {"source", "src", "term", "source term", "مصطلح", "عربي", "arabe", "français", "francais", "fr", "ar", "en"}
_HEADER_WORDS = {"source", "target", "term", "terme", "traduction", "translation", "مصطلح", "ترجمة", "عربي", "arabe",
                 "français", "francais", "english", "anglais", "domain", "domaine", "note", "مجال"}


def _looks_like_header(row: list[str]) -> bool:
    return bool(row) and all(normalize(c) in {normalize(w) for w in _HEADER_WORDS} or not c.strip() for c in row)


def _rows_to_terms(rows, origin: str, cols: tuple[int, int] = (0, 1)):
    rows = [[str(c).strip() if c is not None else "" for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if rows and _looks_like_header(rows[0]):
        rows = rows[1:]
    a, b = cols
    for r in rows:
        if len(r) > max(a, b):
            yield Term(r[a], r[b], domain=r[2] if len(r) > 2 else "", note=r[3] if len(r) > 3 else "", origin=origin)


def _read_text(path: Path) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1256", "cp1252"):
        try:
            return path.read_text(encoding=enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def import_file(path: str | Path, cols: tuple[int, int] = (0, 1)) -> list[Term]:
    """Read one glossary file. `cols` selects the (source, target) columns."""
    p = Path(path)
    suffix, origin = p.suffix.lower(), p.name
    if suffix in {".csv", ".tsv"}:
        text = _read_text(p)
        dialect = csv.excel_tab if suffix == ".tsv" else csv.Sniffer().sniff(text[:2048], delimiters=",;\t|") if text.strip() else csv.excel
        return list(_rows_to_terms(csv.reader(text.splitlines(), dialect), origin, cols))
    if suffix == ".txt":
        rows = [_TXT_SPLIT.split(line, maxsplit=1) for line in _read_text(p).splitlines() if line.strip()]
        return list(_rows_to_terms(rows, origin, cols))
    if suffix == ".json":
        data = json.loads(_read_text(p))
        if isinstance(data, dict):
            data = [{"source": k, "target": v} for k, v in data.items()]
        return [Term(origin=origin, **{k: v for k, v in d.items() if k in Term.__dataclass_fields__ and k != "origin"})
                for d in data]
    if suffix == ".xlsx":
        try:
            import openpyxl
        except ImportError as e:
            raise RuntimeError("pip install openpyxl to import .xlsx glossaries") from e
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        return [t for ws in wb.worksheets for t in _rows_to_terms(ws.iter_rows(values_only=True), origin, cols)]
    if suffix == ".docx":
        try:
            import docx
        except ImportError as e:
            raise RuntimeError("pip install python-docx to import .docx glossaries") from e
        d = docx.Document(str(p))
        return [t for tbl in d.tables
                for t in _rows_to_terms(([c.text for c in row.cells] for row in tbl.rows), origin, cols)]
    raise ValueError(f"Unsupported glossary format: {p.suffix}")


def import_path(path: str | Path, cols: tuple[int, int] = (0, 1)) -> Glossary:
    """Import a file, or every supported file found recursively in a folder."""
    p = Path(path)
    files = sorted(f for f in p.rglob("*") if f.suffix.lower() in SUPPORTED_SUFFIXES) if p.is_dir() else [p]
    g = Glossary()
    for f in files:
        g.extend(import_file(f, cols))
    return g
