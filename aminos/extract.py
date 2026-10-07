"""Extract paragraphs from .docx, .txt, .pdf and images (OCR for scans)."""
from __future__ import annotations

import io
from pathlib import Path
from typing import Callable

# OCR callable: (image bytes, media type) -> transcribed text
OCR = Callable[[bytes, str], str]

IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
               ".webp": "image/webp", ".gif": "image/gif"}


def ocr_claude(model: str = "claude-sonnet-5-5") -> OCR:
    """Transcription-only OCR through Claude vision (handles Arabic/French handwriting-free scans, stamps)."""
    import base64
    import anthropic
    client = anthropic.Anthropic()
    prompt = ("Transcris fidèlement tout le texte de cette page, dans la langue d'origine, sans traduire, "
              "sans commenter. Un paragraphe par ligne vide. Transcris aussi le texte des cachets et sceaux "
              "précédé de [Cachet : ...] et signale logos et photos entre crochets.")

    def ocr(data: bytes, media_type: str) -> str:
        resp = client.messages.create(model=model, max_tokens=4000, messages=[{"role": "user", "content": [
            {"type": "image", "source": {"type": "base64", "media_type": media_type,
                                         "data": base64.b64encode(data).decode()}},
            {"type": "text", "text": prompt}]}])
        return "".join(b.text for b in resp.content if b.type == "text").strip()

    return ocr


def ocr_tesseract(lang: str = "ara+fra+eng") -> OCR:
    """Offline OCR (needs the tesseract binary + pytesseract + Pillow)."""
    import pytesseract
    from PIL import Image

    def ocr(data: bytes, media_type: str) -> str:
        return pytesseract.image_to_string(Image.open(io.BytesIO(data)), lang=lang).strip()

    return ocr


def _split(text: str) -> list[str]:
    blocks = [b.strip() for b in text.split("\n\n") if b.strip()]
    return blocks if len(blocks) > 1 else [l.strip() for l in text.splitlines() if l.strip()]


def _pdf(path: Path, ocr: OCR | None) -> list[str]:
    try:
        import pdfplumber
    except ImportError as e:
        raise RuntimeError("pip install pdfplumber") from e
    out: list[str] = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            text = (page.extract_text() or "").strip()
            if not text and not page.images:
                continue  # blank page
            if len(text) < 20 and page.images:  # scanned page: no usable text layer
                if ocr is None:
                    raise RuntimeError(f"{path.name} p.{page.page_number}: page scannée, activez l'OCR")
                buf = io.BytesIO()
                page.to_image(resolution=200).original.save(buf, format="PNG")
                text = ocr(buf.getvalue(), "image/png")
            out.extend(_split(text))
    return out


def extract_paragraphs(path: str | Path, ocr: OCR | None = None) -> list[str]:
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".docx":
        import docx
        d = docx.Document(str(p))
        paras = [x.text for x in d.paragraphs if x.text.strip()]
        for tbl in d.tables:
            paras += [" | ".join(c.text.strip() for c in row.cells) for row in tbl.rows]
        return paras
    if suffix == ".pdf":
        return _pdf(p, ocr)
    if suffix in IMAGE_TYPES:
        if ocr is None:
            raise RuntimeError("OCR requis pour les images")
        return _split(ocr(p.read_bytes(), IMAGE_TYPES[suffix]))
    return _split(p.read_text(encoding="utf-8"))
