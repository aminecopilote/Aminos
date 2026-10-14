import asyncio
import tempfile
import threading
import time
import unittest
from pathlib import Path

from aminos.glossary import Glossary, Term
from aminos.prompts import DELIMITER
from aminos.tm import TranslationMemory

try:
    from aminos.swarm import HAVE_COORDINATOR, Strategy, run_job
except ImportError:
    HAVE_COORDINATOR = False

AR = "الحكم"


def make_chat(good: bool, delay: float = 0.0, fail: bool = False):
    """Scripted model: translates every segment to 'le jugement' (good) or 'la décision' (bad)."""
    lock, state = threading.Lock(), {"calls": 0}
    answer = "le jugement" if good else "la décision"

    def chat(messages, system):
        with lock:
            state["calls"] += 1
        if fail:
            raise RuntimeError("429 rate limit")
        time.sleep(delay)
        last = messages[-1]["content"]
        n = messages[0]["content"].count(DELIMITER) + 1
        if last.startswith("Retraduis"):
            return f"\n{DELIMITER}\n".join([AR] * n)
        if last.startswith("Compare"):
            return "RAS"
        return f"\n{DELIMITER}\n".join([answer] * n)

    chat.state = state
    return chat


def gloss():
    g = Glossary()
    g.add(Term(AR, "jugement"))
    return g


def run(*a, **k):
    return asyncio.run(run_job(*a, **k))


@unittest.skipUnless(HAVE_COORDINATOR, 'needs agent-coordinator: pip install -e ".[swarm]"')
class SwarmTests(unittest.TestCase):
    def test_clean_first_round(self):
        r = run([AR, AR], "arabe", "français", [make_chat(True)], gloss())
        self.assertTrue(r.ok)
        self.assertEqual(r.paragraphs, ["le jugement"] * 2)
        self.assertEqual([x["round"] for x in r.rounds], [1])

    def test_escalates_to_second_model_until_clean(self):
        bad, good = make_chat(False), make_chat(True)
        r = run([AR], "arabe", "français", [bad, good], gloss())
        self.assertTrue(r.ok, r.remaining_issues)
        self.assertEqual([(x["mode"], x["chat"]) for x in r.rounds], [("normal", 0), ("hard", 0), ("hard", 1)])
        self.assertEqual(r.paragraphs, ["le jugement"])
        self.assertGreater(bad.state["calls"], 0)

    def test_provider_outage_is_survived_by_next_strategy(self):
        r = run([AR], "arabe", "français", [make_chat(True, fail=True), make_chat(True)], gloss())
        self.assertTrue(r.ok)
        self.assertEqual(r.rounds[-1]["chat"], 1)

    def test_exhausted_ladder_reports_failure_not_success(self):
        sealed = []
        r = run([AR], "arabe", "français", [make_chat(False)], gloss(), notarize=sealed.append)
        self.assertFalse(r.ok)
        self.assertEqual(len(r.rounds), 3)
        self.assertTrue(any("jugement" in m for m in r.remaining_issues), r.remaining_issues)
        self.assertEqual(sealed, [])

    def test_only_failing_chunks_are_rerun(self):
        # chunk 1 clean on the first model; chunk 2 only fixed by the second model
        calls = {"n": 0}

        def mixed(messages, system):
            calls["n"] += 1
            n = messages[0]["content"].count(DELIMITER) + 1
            last = messages[-1]["content"]
            if last.startswith(("Retraduis", "Compare")):
                return "RAS" if last.startswith("Compare") else AR
            wanted_good = "ONE" in messages[0]["content"]
            return f"\n{DELIMITER}\n".join(["le jugement" if wanted_good else "la décision"] * n)

        r = run([AR + " ONE", AR + " TWO"], "arabe", "français", [mixed, make_chat(True)], gloss(), max_chars=3)
        self.assertTrue(r.ok, r.remaining_issues)
        self.assertEqual(r.rounds[0]["chunks_run"], 2)
        self.assertEqual(r.rounds[0]["chunks_clean"], 1)
        self.assertEqual(r.rounds[1]["chunks_run"], 1)

    def test_translation_memory_reuse_and_learning(self):
        tm = TranslationMemory()
        tm.add("نص قديم", "ancien texte", "arabe", "français")
        chat = make_chat(True)
        r = run(["نص قديم", AR], "arabe", "français", [chat], gloss(), tm=tm)
        self.assertEqual((r.paragraphs, r.reused_from_tm), (["ancien texte", "le jugement"], 1))
        self.assertEqual(tm.lookup(AR, "arabe", "français")[0].target, "le jugement")  # learned
        self.assertEqual(chat.state["calls"], 1)

    def test_translators_really_run_in_parallel(self):
        paras = [AR] * 6  # no digits: the number check would rightly flag a missing "3"
        t0 = time.time()
        r = run(paras, "arabe", "français", [make_chat(True, delay=0.3)], gloss(),
                ladder=(Strategy("fast", 0),), translators=3, max_chars=3)
        elapsed = time.time() - t0
        self.assertTrue(r.ok)
        self.assertEqual(r.rounds[0]["chunks_run"], 6)
        self.assertLess(elapsed, 1.3, f"serial would take ~1.8 s, took {elapsed:.2f}")

    def test_docx_and_notarisation_only_when_clean(self):
        try:
            import docx  # noqa: F401
            from aminos.notary import evidence
            from aminos.notary.pki import create_test_pki
        except ImportError:
            self.skipTest("needs python-docx and the notary extra")
        d = Path(tempfile.mkdtemp())
        pki = create_test_pki(d / "pki")

        def notarize(path):
            ev = evidence.notarize(path, d, pki)
            return evidence.check(path, d, pki["ca"], ev)

        r = run([AR], "arabe", "français", [make_chat(True)], gloss(), docx_out=d / "out.docx", notarize=notarize)
        self.assertTrue(r.ok)
        self.assertTrue((d / "out.docx").exists())
        self.assertTrue(r.notary["ok"], r.notary)


if __name__ == "__main__":
    unittest.main()
