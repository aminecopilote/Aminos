import tempfile
import unittest
import zipfile
from pathlib import Path

from aminos import certify as c

PARAS = ["**Royaume du Maroc**", "Article 1 : jugement rendu le 2026-08-20 à **Rabat**.", "Article 2 : le pourvoi est rejeté."]


def setup(d):
    d = Path(d)
    key = d / "k.key"
    c.make_key(key)
    return d, key, d / "reg.csv"


class DeclarationTests(unittest.TestCase):
    def test_single_language_and_no_removed_mention(self):
        for t, city in (("français", "Rabat"), ("arabe", "الرباط"), ("anglais", "Rabat")):
            text = c.declaration("arabe" if t != "arabe" else "français", t, "00012.2026", "2026-09-01")
            self.assertIn(city, text)
            self.assertNotIn("ONG", text)
            self.assertNotIn("NGO", text)
        self.assertIn("1er septembre 2026", c.declaration("arabe", "français", date="2026-09-01"))
        self.assertIn("شتنبر", c.declaration("français", "arabe", date="2026-09-01"))
        self.assertIn("traduction de l'arabe vers le français", c.declaration("arabe", "français"))
        self.assertIn("traduction de l'anglais vers le français", c.declaration("anglais", "français"))
        self.assertIn("from Arabic into English", c.declaration("arabe", "anglais"))

    def test_arabic_declaration_has_no_latin_letters(self):
        text = c.declaration("français", "arabe", "00012.2026", "2026-09-01")
        self.assertFalse([ch for ch in text if ch.isascii() and ch.isalpha()])


class RegistryTests(unittest.TestCase):
    def test_numbering(self):
        with tempfile.TemporaryDirectory() as d:
            reg = Path(d) / "r.csv"
            self.assertEqual(c.next_ref(reg, 2026), "00001.2026")
            c.upsert(reg, ref="00041.2025", date="2025-01-01", code="x", seal="", status="scellé")
            self.assertEqual(c.next_ref(reg, 2026), "00042.2026")


class CertificationFlowTests(unittest.TestCase):
    def build(self, d, target="français"):
        d, key, reg = setup(d)
        r = c.certify(PARAS, "arabe", target, key, reg, d / "archive", date="2026-09-01")
        return d, key, reg, r

    def test_code_stable_with_and_without_block_and_verifies(self):
        with tempfile.TemporaryDirectory() as d:
            d, key, reg, r = self.build(d)
            self.assertTrue(c.verify(r["path"], key, reg, r["ref"]))
            bare = d / "bare.docx"
            c.build_certified_docx(PARAS, "arabe", "français", r["ref"], "2026-09-01", bare)
            self.assertEqual(c.body_text(bare), c.body_text(r["path"]))

    def test_tampering_breaks_verification(self):
        with tempfile.TemporaryDirectory() as d:
            d, key, reg, r = self.build(d)
            tampered = d / "t.docx"
            c.build_certified_docx(PARAS[:-1] + ["Article 2 : le pourvoi est accepté."], "arabe", "français",
                                   r["ref"], "2026-09-01", tampered)
            self.assertFalse(c.verify(tampered, key, reg, r["ref"]))

    def test_qr_roundtrip(self):
        payload = c.qr_payload("00001.2026", "2026-09-01", "AB" * 32)
        self.assertEqual(c.read_qr(c.make_qr_png(payload)), payload)

    def test_no_header_footer_text_and_banner_in_body(self):
        import docx
        with tempfile.TemporaryDirectory() as d:
            d, key, reg, r = self.build(d)
            doc = docx.Document(r["path"])
            self.assertEqual(doc.paragraphs[0].text, "TRADUCTION CERTIFIÉE CONFORME")
            for s in doc.sections:
                self.assertEqual([p.text for p in s.header.paragraphs if p.text], [])
                self.assertEqual([p.text for p in s.footer.paragraphs if p.text], [])
            self.assertEqual(len(doc.tables), 1)
            self.assertNotIn("ONG", "\n".join(p.text for p in doc.paragraphs))

    def test_one_page_when_libreoffice_available(self):
        with tempfile.TemporaryDirectory() as d:
            d, key, reg, r = self.build(d)
            n = c.page_count(r["path"])
            if n is not None:
                self.assertEqual(n, 1)

    def test_arabic_layout_uses_left_for_rtl_start(self):
        import docx
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        with tempfile.TemporaryDirectory() as d:
            d, key, reg, r = self.build(d, "arabe")
            doc = docx.Document(r["path"])
            body = doc.paragraphs[1]
            self.assertEqual(body.alignment, WD_ALIGN_PARAGRAPH.LEFT)
            self.assertIn("w:bidi", body._p.xml)
            self.assertIn(c.LRI + "2026-08-20" + c.PDI, "".join(p.text for p in doc.paragraphs))

    def test_seal_requires_validation(self):
        with tempfile.TemporaryDirectory() as d:
            d, key, reg, r = self.build(d)
            ev = d / "ev.md"
            good = c.Evaluation({k: 5 for k in c.CRITERIA}, ["Nom illisible p.2"])
            c.write_evaluation(ev, r["ref"], good)
            with self.assertRaises(PermissionError):   # not validated by Alami yet
                c.seal(r["path"], key, reg, r["ref"], ev)
            ev.write_text(ev.read_text(encoding="utf-8").replace("Validation Alami : non", "Validation Alami : oui"),
                          encoding="utf-8")
            s = c.seal(r["path"], key, reg, r["ref"], ev)
            self.assertEqual(len(s), 64)
            self.assertEqual(c.read_registry(reg)[0]["status"], "scellé")
            with self.assertRaises(ValueError):          # cannot re-certify a sealed reference
                c.certify(PARAS, "arabe", "français", key, reg, d / "archive", ref=r["ref"], date="2026-09-01")

    def test_threshold(self):
        weak = {k: 5 for k in c.CRITERIA} | {"fidelite": 4}
        self.assertFalse(c.Evaluation(weak, []).passes())
        self.assertFalse(c.Evaluation({k: 5 for k in c.CRITERIA} | {"formats": 3}, []).passes())
        with tempfile.TemporaryDirectory() as d:
            ev = Path(d) / "e.md"
            c.write_evaluation(ev, "00001.2026", c.Evaluation(weak, []))
            ev.write_text(ev.read_text(encoding="utf-8").replace(": non", ": oui"), encoding="utf-8")
            self.assertFalse(c.evaluation_validated(ev))   # validated by hand but below threshold

    def test_key_never_overwritten(self):
        with tempfile.TemporaryDirectory() as d:
            _, key, _ = setup(d)
            with self.assertRaises(FileExistsError):
                c.make_key(key)


if __name__ == "__main__":
    unittest.main()
