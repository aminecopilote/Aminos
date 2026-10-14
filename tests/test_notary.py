import json
import shutil
import tempfile
import unittest
from pathlib import Path

from aminos.notary import Ledger, LedgerError, sha256_bytes, sha256_file

try:
    import docx
    from fpdf import FPDF
    import pyhanko  # noqa: F401
    from aminos.notary import evidence, seal
    from aminos.notary.pki import create_test_pki
    HAVE_CRYPTO = shutil.which("openssl") is not None
except ImportError:  # optional extra: pip install -e ".[notary]"
    HAVE_CRYPTO = False


class LedgerTests(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())
        self.ledger = Ledger(self.dir / "l.jsonl")
        for i in range(3):
            self.ledger.append(sha256_bytes(str(i).encode()), f"doc{i}", now=1000 + i)

    def test_chain_valid_and_head_moves(self):
        self.assertEqual(self.ledger.verify(), 3)
        self.assertEqual(self.ledger.head(), self.ledger.entries()[-1]["hash"])

    def test_edit_is_detected(self):
        lines = (self.dir / "l.jsonl").read_text().splitlines()
        lines[1] = lines[1].replace("doc1", "docX")
        (self.dir / "l.jsonl").write_text("\n".join(lines) + "\n")
        with self.assertRaisesRegex(LedgerError, "entry 2: content modified"):
            self.ledger.verify()

    def test_deleted_entry_is_detected(self):
        lines = (self.dir / "l.jsonl").read_text().splitlines()
        (self.dir / "l.jsonl").write_text("\n".join([lines[0], lines[2]]) + "\n")
        with self.assertRaisesRegex(LedgerError, "entry 2"):
            self.ledger.verify()

    def test_cannot_extend_broken_chain(self):
        lines = (self.dir / "l.jsonl").read_text().splitlines()
        (self.dir / "l.jsonl").write_text("\n".join(lines[1:]) + "\n")
        with self.assertRaises(LedgerError):
            self.ledger.append("a" * 64)


@unittest.skipUnless(HAVE_CRYPTO, "needs cryptography, pyhanko, fpdf2, python-docx and openssl")
class SealTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dir = Path(tempfile.mkdtemp())
        cls.pki = create_test_pki(cls.dir / "pki")
        cls.other = create_test_pki(cls.dir / "other")  # unrelated CA

    def fresh(self, kind):
        d = Path(tempfile.mkdtemp())
        if kind == "docx":
            x = docx.Document(); x.add_paragraph("Acte fictif n° 1"); x.save(d / "doc.docx")
            return d, d / "doc.docx"
        pdf = FPDF(); pdf.add_page(); pdf.set_font("helvetica", size=12); pdf.cell(0, 10, "Document fictif")
        pdf.output(d / "doc.pdf")
        return d, d / "doc.pdf"

    def test_docx_roundtrip_and_independent_openssl_check(self):
        d, doc = self.fresh("docx")
        ev = evidence.notarize(doc, d, self.pki)
        self.assertTrue(evidence.check(doc, d, self.pki["ca"], ev)["ok"])

    def test_docx_modified_after_sealing(self):
        d, doc = self.fresh("docx")
        ev = evidence.notarize(doc, d, self.pki)
        doc.write_bytes(doc.read_bytes() + b"\0")
        r = evidence.check(doc, d, self.pki["ca"], ev)
        self.assertFalse(r["ok"])
        self.assertFalse(r["checks"]["sha256_matches_evidence"])
        self.assertFalse(r["checks"]["seal_valid"])

    def test_seal_from_untrusted_ca_is_rejected(self):
        d, doc = self.fresh("docx")
        ev = evidence.notarize(doc, d, self.pki)
        self.assertFalse(evidence.check(doc, d, self.other["ca"], ev)["checks"]["seal_valid"])

    def test_pdf_roundtrip(self):
        d, doc = self.fresh("pdf")
        ev = evidence.notarize(doc, d, self.pki)
        r = evidence.check(doc, d, self.pki["ca"], ev)
        self.assertTrue(r["ok"], r)
        info = seal.verify_pdf(d / "doc.sealed.pdf", self.pki["ca"])
        self.assertEqual((info["intact"], info["valid"], info["trusted"]), (True, True, True))
        self.assertIn("TEST", info["signer"].upper() + "TEST")

    def test_pdf_tampered_after_sealing(self):
        d, doc = self.fresh("pdf")
        evidence.notarize(doc, d, self.pki)
        sealed = d / "doc.sealed.pdf"
        raw = bytearray(sealed.read_bytes())
        i = raw.index(b"Document fictif") if b"Document fictif" in raw else 200
        raw[i] ^= 1
        sealed.write_bytes(bytes(raw))
        info = seal.verify_pdf(sealed, self.pki["ca"])
        self.assertFalse(info["intact"] and info["valid"])

    def test_pdf_sealed_by_unknown_ca_not_trusted(self):
        d, doc = self.fresh("pdf")
        evidence.notarize(doc, d, self.pki)
        self.assertFalse(seal.verify_pdf(d / "doc.sealed.pdf", self.other["ca"])["trusted"])

    def test_ledger_tamper_surfaces_in_check(self):
        d, doc = self.fresh("docx")
        ev = evidence.notarize(doc, d, self.pki)
        (d / "ledger.jsonl").write_text((d / "ledger.jsonl").read_text().replace(sha256_file(doc), "f" * 64))
        r = evidence.check(doc, d, self.pki["ca"], ev)
        self.assertFalse(r["checks"]["ledger_chain_intact"])
        self.assertFalse(r["ok"])

    def test_private_key_permissions_and_ca_has_no_key(self):
        self.assertEqual((self.pki["key"].stat().st_mode & 0o777), 0o600)
        self.assertFalse(list((self.dir / "pki").glob("ca.key")))


if __name__ == "__main__":
    unittest.main()
